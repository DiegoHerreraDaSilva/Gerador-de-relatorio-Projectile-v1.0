"""Tabela do chat analítico → .xlsx (`POST /analytics/chat/export`).

Não consulta nada: formata a tabela que o navegador já exibiu. Números
continuam números no Excel (dá pra somar/filtrar), com o formato de cada
coluna (`column_types`: hours | percent | count | ms | text).

Não é o caminho de geração de relatório — aqui não há template, então
`Workbook.save()` é seguro (a regra "nunca openpyxl.save()" do CLAUDE.md é
sobre `generator.py`, que precisa preservar desenhos do template)."""
from __future__ import annotations

import io
import re
import unicodedata

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .schemas import ExportRequest

_FORMATS = {
    "hours": '#,##0.00" h"',
    "percent": '0.0"%"',
    "count": "#,##0",
    "ms": '#,##0" ms"',
}
_HEADER_FILL = PatternFill("solid", fgColor="0F8F9C")
_TOTAL_FILL = PatternFill("solid", fgColor="E6F2F3")
_THIN = Side(style="thin", color="C9D3D8")


def _sheet_title(title: str) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", title).strip() or "Tabela"
    return cleaned[:31]


def filename_for(title: str) -> str:
    plain = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^A-Za-z0-9]+", "-", plain).strip("-").lower()[:80] or "tabela"
    return f"{slug}.xlsx"


def _write_cell(ws, row: int, column: int, value, kind: str | None):
    cell = ws.cell(row=row, column=column)
    if isinstance(value, str):
        cell.value = value
        # texto nunca vira fórmula ("=HYPERLINK(...)" num nome de pacote)
        cell.data_type = "s"
    else:
        cell.value = value
        if value is not None and kind in _FORMATS:
            cell.number_format = _FORMATS[kind]
    return cell


def build_xlsx(payload: ExportRequest) -> bytes:
    columns = payload.columns
    kinds = list(payload.column_types or [])
    kinds += ["text"] * (len(columns) - len(kinds))
    wb = Workbook()
    ws = wb.active
    ws.title = _sheet_title(payload.title)

    ws.cell(row=1, column=1, value=payload.title).font = Font(bold=True, size=13)
    header_row = 3
    for index, name in enumerate(columns, start=1):
        cell = _write_cell(ws, header_row, index, name, "text")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=_THIN)

    widths = [len(str(name)) for name in columns]
    current = header_row
    for values in payload.rows:
        current += 1
        for index, value in enumerate(values[: len(columns)], start=1):
            _write_cell(ws, current, index, value, kinds[index - 1])
            widths[index - 1] = max(widths[index - 1], len(str(value)) if value is not None else 0)
    if payload.totals:
        current += 1
        for index, value in enumerate(payload.totals[: len(columns)], start=1):
            cell = _write_cell(ws, current, index, value, kinds[index - 1])
            cell.font = Font(bold=True)
            cell.fill = _TOTAL_FILL
            cell.border = Border(top=_THIN)

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=2)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{max(current, header_row)}"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
