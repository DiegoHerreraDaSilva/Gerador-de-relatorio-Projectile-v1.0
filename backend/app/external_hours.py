"""Horas externas: planilha própria de colaboradores que NÃO apontam no Projectile.

Não é o export do Projectile (esse é o `parser.py`): é um modelo simples, uma linha
por lançamento — Data, Colaborador, Projeto, Pacote de Trabalho, Descrição, Horas. O
parser só LÊ e valida; quem decide pra onde cada linha vai no relatório aberto é o
frontend (`utils/externalHours.ts`), que conhece os pacotes e grupos da tela.

Função pura, sem banco nem rede, pra testar com planilhas montadas em memória."""

from __future__ import annotations

import io
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils.datetime import from_excel

from .parser import _parse_hs_value

MAX_ROWS = 5000
MAX_HOURS_PER_ROW = 24.0
MAX_DESCRIPTION = 500  # o payload de geração aceita 2000; aqui sobra espaço pro "Colaborador – "
MAX_NAME = 300
MAX_COLLABORATOR = 120

# coluna do modelo -> (rótulo no modelo, apelidos aceitos no cabeçalho, já normalizados)
COLUMNS: dict[str, tuple[str, tuple[str, ...]]] = {
    "data": ("Data", ("data", "dia", "data do lancamento")),
    "colaborador": ("Colaborador", ("colaborador", "colaboradora", "pessoa", "nome", "funcionario", "profissional")),
    "projeto": ("Projeto", ("projeto",)),
    "pacote": ("Pacote de Trabalho", ("pacote de trabalho", "pacote", "pacote trabalho")),
    "descricao": ("Descrição", ("descricao", "atividade", "observacao", "descricao da atividade")),
    "horas": ("Horas", ("horas", "hs", "horas trabalhadas", "qtd horas", "total de horas")),
}

_CONTROL = re.compile(r"[\x00-\x08\x0b-\x1f\x7f]")


@dataclass(frozen=True)
class ExternalRow:
    row: int  # linha na planilha (1-based)
    date: date
    collaborator: str
    project: str
    package: str
    description: str
    hours: float

    def as_dict(self) -> dict:
        return {
            "row": self.row,
            "date": self.date.isoformat(),
            "collaborator": self.collaborator,
            "project": self.project,
            "package": self.package,
            "description": self.description,
            "hours": self.hours,
        }


@dataclass(frozen=True)
class ExternalIssue:
    row: int
    reason: str  # campo_vazio | data_invalida | horas_invalidas | texto_longo
    message: str

    def as_dict(self) -> dict:
        return {"row": self.row, "reason": self.reason, "message": self.message}


def _norm(value) -> str:
    """Sem acento, sem caixa, só letras/números/espaço — pra casar cabeçalho."""
    text = unicodedata.normalize("NFD", str(value or "").lower())
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", " ", text)).strip()


def _clean(value) -> str:
    return re.sub(r"\s+", " ", _CONTROL.sub("", str(value if value is not None else ""))).strip()


def _cell(values: tuple, columns: dict[str, int], key: str):
    position = columns[key]
    return values[position] if position < len(values) else None


def _find_header(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    """Primeira linha (entre as 30 primeiras) que traz as seis colunas. Se nenhuma
    traz, o erro lista o que faltou na linha mais parecida."""
    best_missing: list[str] | None = None
    for index, values in enumerate(rows[:30]):
        found: dict[str, int] = {}
        for position, cell in enumerate(values):
            name = _norm(cell)
            if not name:
                continue
            for key, (_, aliases) in COLUMNS.items():
                if key not in found and name in aliases:
                    found[key] = position
        missing = [COLUMNS[key][0] for key in COLUMNS if key not in found]
        if not missing:
            return index, found
        if found and (best_missing is None or len(missing) < len(best_missing)):
            best_missing = missing
    if best_missing is None:
        raise ValueError(
            "Não encontrei o cabeçalho do modelo (Data, Colaborador, Projeto, Pacote de Trabalho, Descrição, Horas). Baixe o modelo e preencha a partir dele."
        )
    raise ValueError(f"Planilha fora do modelo: faltam as colunas {', '.join(best_missing)}. Baixe o modelo e preencha a partir dele.")


def _parse_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and 20000 <= value <= 80000:
        try:
            return from_excel(value).date()  # data digitada sem formato de data: número de série do Excel
        except (ValueError, OverflowError):
            return None
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
    return None


def parse_external_hours(path_or_file) -> tuple[list[ExternalRow], list[ExternalIssue]]:
    """Lê a primeira aba. Linha totalmente em branco, ou só com número (subtotal/rodapé),
    é ignorada em silêncio; linha com algum campo de texto mas incompleta ou inválida
    vira aviso e fica de fora."""
    workbook = openpyxl.load_workbook(path_or_file, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        all_rows = [tuple(values) for values in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()

    header_index, columns = _find_header(all_rows)
    rows: list[ExternalRow] = []
    issues: list[ExternalIssue] = []
    seen = 0
    for offset, values in enumerate(all_rows[header_index + 1 :], start=header_index + 2):
        text_fields = {key: _clean(_cell(values, columns, key)) for key in ("colaborador", "projeto", "pacote", "descricao")}
        raw_date, raw_hours = _cell(values, columns, "data"), _cell(values, columns, "horas")
        if not any(text_fields.values()):
            continue  # em branco ou subtotal/rodapé (só data/número)
        seen += 1
        if seen > MAX_ROWS:
            raise ValueError(f"A planilha tem mais de {MAX_ROWS} linhas de lançamento. Divida em arquivos menores.")

        empty = [COLUMNS[key][0] for key, text in text_fields.items() if not text]
        if empty:
            issues.append(ExternalIssue(offset, "campo_vazio", f"Linha {offset}: falta preencher {', '.join(empty)}."))
            continue
        day = _parse_date(raw_date)
        if day is None:
            issues.append(ExternalIssue(offset, "data_invalida", f"Linha {offset}: data inválida ({_clean(raw_date) or 'vazia'}). Use dd/mm/aaaa."))
            continue
        try:
            hours = _parse_hs_value(raw_hours)
        except ValueError:
            issues.append(ExternalIssue(offset, "horas_invalidas", f"Linha {offset}: horas inválidas ({_clean(raw_hours) or 'vazio'})."))
            continue
        if not 0 < hours <= MAX_HOURS_PER_ROW:  # também recusa NaN
            issues.append(ExternalIssue(offset, "horas_invalidas", f"Linha {offset}: horas fora do limite (0 a {MAX_HOURS_PER_ROW:g} por linha): {hours:g}."))
            continue
        if (
            len(text_fields["descricao"]) > MAX_DESCRIPTION
            or len(text_fields["colaborador"]) > MAX_COLLABORATOR
            or max(len(text_fields["projeto"]), len(text_fields["pacote"])) > MAX_NAME
        ):
            issues.append(ExternalIssue(offset, "texto_longo", f"Linha {offset}: texto longo demais (descrição até {MAX_DESCRIPTION} caracteres)."))
            continue
        rows.append(
            ExternalRow(
                row=offset,
                date=day,
                collaborator=text_fields["colaborador"],
                project=text_fields["projeto"],
                package=text_fields["pacote"],
                description=text_fields["descricao"],
                hours=round(hours, 3),
            )
        )
    return rows, issues


def build_template() -> bytes:
    """O modelo pra baixar: aba de lançamentos só com o cabeçalho (exemplo preenchido ali
    seria importado se o usuário esquecesse de apagar) e uma aba de instruções com um
    exemplo. Não é o caminho de geração de relatório — `Workbook.save()` é seguro aqui."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Horas externas"
    sheet.append([label for label, _ in COLUMNS.values()])
    for letter, width in zip("ABCDEF", (14, 28, 36, 34, 60, 10), strict=True):
        sheet.column_dimensions[letter].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"
    for row in sheet.iter_rows(min_row=2, max_row=MAX_ROWS + 1, min_col=1, max_col=1):
        row[0].number_format = "dd/mm/yyyy"

    help_sheet = workbook.create_sheet("Instruções")
    lines = [
        ["Como preencher"],
        [],
        ["1. Uma linha por lançamento, na aba 'Horas externas'. Não mude os nomes das colunas."],
        [
            "2. Projeto e Pacote de Trabalho devem ter o mesmo nome que aparece no relatório aberto; linha sem correspondência fica pendente e você escolhe o destino na tela."
        ],
        ["3. Data no formato dd/mm/aaaa, dentro do período do relatório."],
        [f"4. Horas em número (2,5 ou 2.5), de 0 a {MAX_HOURS_PER_ROW:g} por linha. Máximo de {MAX_ROWS} linhas."],
        ["5. No relatório, a descrição sai como 'Colaborador – Descrição'."],
        [],
        ["Exemplo (não está na aba de lançamentos, de propósito)"],
        [label for label, _ in COLUMNS.values()],
        ["12/08/2026", "Fulano de Tal", "Projeto Exemplo 08.2026", "Pacote de Trabalho Exemplo", "Revisão do desenho", 3.5],
    ]
    for line in lines:
        help_sheet.append(line)
    help_sheet.column_dimensions["A"].width = 110
    help_sheet["A1"].font = Font(bold=True, size=13)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
