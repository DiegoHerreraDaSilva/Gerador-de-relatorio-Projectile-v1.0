"""Relatório gerado por QUALQUER meio (download, envio, aprovação/envio da
geração automática) nunca mostra Bruto/Performance — decisão do usuário
(2026-09-28). O pedido pode vir com `include_performance: true` (guia
antiga, chamada direta na API): o arquivo sai igual, sem as colunas. A
performance continua no CÁLCULO das horas."""
from __future__ import annotations

import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader

from backend.app.api.dependencies import require_session
from backend.app.main import app
from backend.app.services import report_persistence
from backend.app.services.report_files import GeneratePayload, build_report_file

_USER = {"name": "Fulano", "login": "fulano", "email": "f@x", "employee_id": "1", "filiale": None}


def _payload(formats: list[str]) -> dict:
    return {
        "packages": [{
            "header": {
                "project_code": "SE.26.001", "project_name": "Projeto", "location_date": "Santo André, 01.09.2026",
                "month_label": "Agosto/2026", "signer1_name": "A", "signer1_company": "Schwaben Engineering",
                "signer2_name": "B", "signer2_company": "Cliente",
            },
            "groups": [{"name": "Grupo A", "performance": 1.1,
                        "activities": [{"description": "Ativ 1", "hours": 10.0}]}],
        }],
        "formats": formats,
        "include_performance": True,   # pedido explícito — tem que ser ignorado
    }


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(report_persistence, "begin_generation", lambda *a, **k: None)
    app.dependency_overrides[require_session] = lambda: _USER
    yield TestClient(app)
    app.dependency_overrides.pop(require_session, None)


def _xlsx_values(data: bytes) -> list:
    ws = load_workbook(io.BytesIO(data), data_only=False).active
    return [c.value for row in ws.iter_rows() for c in row if c.value is not None]


def test_download_pedindo_performance_sai_sem_as_colunas(client):
    response = client.post("/generate", json=_payload(["xlsx", "pdf"]))
    assert response.status_code == 200, response.text
    files = {n: zipfile.ZipFile(io.BytesIO(response.content)).read(n)
             for n in zipfile.ZipFile(io.BytesIO(response.content)).namelist()}
    xlsx = next(v for n, v in files.items() if n.endswith(".xlsx"))
    pdf = next(v for n, v in files.items() if n.endswith(".pdf"))
    values = _xlsx_values(xlsx)
    assert "Bruto" not in values and "Performance" not in values
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf)).pages)
    assert "Bruto" not in text and "Performance" not in text
    # a performance continua no cálculo: 10 h × 1,1 = 11 h
    assert "11 h" in text


def test_porta_unica_de_saida_nao_aceita_performance(tmp_path):
    """`build_report_file` é por onde /generate, /send-report e a geração
    automática geram — nem tem mais como pedir as colunas."""
    import inspect

    assert "include_performance" not in inspect.signature(build_report_file).parameters
    pkg = GeneratePayload.model_validate(_payload(["xlsx"])).packages[0]
    path = tmp_path / "r.xlsx"
    build_report_file(pkg, str(path), "xlsx")
    values = _xlsx_values(path.read_bytes())
    assert "Bruto" not in values and "Performance" not in values
