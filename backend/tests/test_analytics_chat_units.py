"""Peças do chat analítico isoladas: cliente do Jev (HTTP falso), trava de
números (grounding), períodos, visualização e o SQL do reports_db."""
from __future__ import annotations

from datetime import date
from types import SimpleNamespace

import pytest
import requests
from fastapi.testclient import TestClient

from backend.app import management
from backend.app.analytics import grounding, periods, visualization
from backend.app.analytics.query_engine import Filters, QueryResult
from backend.app.integrations import jev
from backend.app.main import app, require_session
from backend.app.repositories import report_analytics_repository

# --- Jev --------------------------------------------------------------------


def _settings(key="k", openrouter=""):
    return SimpleNamespace(
        openrouter_api_key=openrouter, typesafe_api_key=key, jev_model="jev-latest", jev_timeout_seconds=3.0,
    )


class _Response:
    def __init__(self, payload, status=200):
        self._payload, self.status_code = payload, status

    def raise_for_status(self):
        if self.status_code >= 400:
            raise requests.HTTPError(f"HTTP {self.status_code}")

    def json(self):
        return self._payload


def test_jev_monta_a_requisicao_e_le_choice_e_noul(monkeypatch):
    sent = {}

    def fake_post(url, headers, json, timeout):
        sent.update(url=url, headers=headers, json=json, timeout=timeout)
        return _Response({"model": "jev-1.13.0", "answers": {
            "intent": {"type": "choice", "choice": "total_hours", "confidence": 0.93,
                       "probabilities": {"total_hours": 0.95, "none": 0.05}},
            "follow_up": {"type": "noul", "noul": 0.9},
            "extra_nao_pedida": {"type": "choice", "choice": "x", "confidence": 1.0},
        }})

    monkeypatch.setattr(jev, "get_settings", lambda: _settings())
    monkeypatch.setattr(jev.requests, "post", fake_post)
    questions = {"intent": jev.choice("q", {"total_hours": "a", "none": "b"}), "follow_up": jev.noul("q")}

    answers = jev.ask("estado", questions)

    assert sent["url"] == jev.TYPESAFE_URL and sent["headers"]["Authorization"] == "Bearer k"
    assert sent["json"] == {"state": "estado", "model": "jev-latest", "questions": questions}
    assert answers["intent"].choice == "total_hours" and answers["intent"].confidence == 0.93
    assert answers["follow_up"].choice == "yes" and answers["follow_up"].confidence == pytest.approx(0.8)
    assert "extra_nao_pedida" not in answers


def test_chave_do_openrouter_tem_prioridade(monkeypatch):
    sent = {}

    def fake_post(url, headers, json, timeout):
        sent.update(url=url, headers=headers)
        return _Response({"answers": {"x": {"type": "noul", "noul": 0.9}}})

    monkeypatch.setattr(jev, "get_settings", lambda: _settings(key="ts", openrouter="or"))
    monkeypatch.setattr(jev.requests, "post", fake_post)
    jev.ask("estado", {"x": jev.noul("q")})
    assert sent["url"] == jev.OPENROUTER_URL and sent["headers"]["Authorization"] == "Bearer or"


def test_jev_sem_chave_nao_faz_chamada(monkeypatch):
    monkeypatch.setattr(jev, "get_settings", lambda: _settings(key=""))
    monkeypatch.setattr(jev.requests, "post", lambda *a, **k: pytest.fail("não deveria chamar"))
    assert jev.is_enabled() is False
    with pytest.raises(jev.ClassifierUnavailableError):
        jev.ask("estado", {"x": jev.noul("q")})


@pytest.mark.parametrize("response", [_Response({}, status=500), _Response({"sem_answers": 1})])
def test_jev_com_erro_vira_indisponivel(monkeypatch, response):
    monkeypatch.setattr(jev, "get_settings", lambda: _settings())
    monkeypatch.setattr(jev.requests, "post", lambda *a, **k: response)
    with pytest.raises(jev.ClassifierUnavailableError):
        jev.ask("estado", {"x": jev.noul("q")})


def test_choice_recusa_mais_de_255_opcoes():
    with pytest.raises(ValueError):
        jev.choice("q", {str(i): "x" for i in range(256)})


# --- grounding ----------------------------------------------------------------


def test_le_numeros_no_formato_brasileiro():
    assert grounding.numbers_in("1.234,5 h, 12,5% e 7 relatórios em 2026") == [1234.5, 12.5, 7.0, 2026.0]


def test_numero_dos_dados_passa_e_numero_inventado_nao():
    allowed = grounding.allowed_numbers({"total": 1234.5, "rows": [{"value": 76.923}], "label": "setembro/2026"})
    assert grounding.is_grounded("Foram 1.234,5 h, 76,9% em setembro de 2026, top 3.", allowed)
    assert not grounding.is_grounded("Foram 1.500 h.", allowed)


# --- períodos -----------------------------------------------------------------

_TODAY = date(2026, 9, 24)


def test_mes_corrente_vai_ate_o_fim_do_mes_como_no_painel():
    """O Painel conta apontamento com data futura no mês corrente."""
    period = periods.resolve_period("2026-09", None, _TODAY, 24)
    assert (period.start, period.end, period.label) == (date(2026, 9, 1), date(2026, 9, 30), "setembro/2026")


def test_periodos_relativos():
    last = periods.resolve_period(None, "last_month", _TODAY, 24)
    assert (last.start, last.end) == (date(2026, 8, 1), date(2026, 8, 31))
    three = periods.resolve_period(None, "last_3_months", _TODAY, 24)
    assert (three.start, three.label) == (date(2026, 7, 1), "julho/2026 a setembro/2026")
    assert three.phrase == "de julho/2026 a setembro/2026"
    assert last.phrase == "em agosto/2026"


def test_intervalo_de_meses_e_ordem_invertida():
    period = periods.resolve_period("2026-02", None, _TODAY, 24, month_end="2026-08")
    assert (period.start, period.end, period.label) == (date(2026, 2, 1), date(2026, 8, 31), "fevereiro/2026 a agosto/2026")
    assert periods.resolve_period("2026-08", None, _TODAY, 24, month_end="2026-02") == period
    assert periods.resolve_period("2026-09", None, _TODAY, 24, month_end="2026-09").label == "setembro/2026"


def test_anos_citados_ignoram_numero_colado_em_texto():
    assert periods.years_mentioned("horas de 2024 a 2025 no PP2030 e 08.2026") == ["2024", "2025", "2026"]
    assert periods.years_mentioned("horas em agosto") == []


def test_detecta_nome_de_mes_com_ou_sem_acento_e_abreviado():
    assert periods.mentions_month("horas em Março de 2025")
    assert periods.mentions_month("horas em marco")
    assert periods.mentions_month("ago/2025")
    assert not periods.mentions_month("horas por cliente em 2026")
    assert not periods.mentions_month("horas da Mercedes")  # "mer" não é mês


def test_mes_fora_da_janela_e_recusado():
    with pytest.raises(ValueError):
        periods.resolve_period("2023-01", None, _TODAY, 24)


def test_opcoes_de_mes_cobrem_a_janela_e_viram_a_virada_de_ano():
    options = periods.month_options(date(2026, 1, 10), 3)
    assert list(options) == ["2026-01", "2025-12", "2025-11"]
    assert "janeiro de 2026" in options["2026-01"]


# --- visualização ------------------------------------------------------------


def _result(intent, dimension, rows, unit="hours", total=None):
    filters = Filters(periods.resolve_period("2026-09", None, _TODAY, 24))
    return QueryResult(intent=intent, source="projectile", unit=unit, dimension=dimension,
                       filters=filters, total=total, rows=rows)


def test_serie_mensal_vira_linha():
    result = _result("hours_by_competence", "competence",
                     [{"label": "agosto/2026", "value": 5.0}, {"label": "setembro/2026", "value": 13.0}])
    assert visualization.build_visualizations(result)[0]["type"] == "line"


def test_ranking_longo_mostra_top_15_e_tabela_completa():
    rows = [{"label": f"Cliente {i:02d}", "value": float(100 - i)} for i in range(40)]
    result = _result("hours_by_client", "client", rows)
    viz = visualization.build_visualizations(result)[0]
    assert viz["type"] == "horizontal_bar" and len(viz["categories"]) == 15
    assert "top 15 de 40" in viz["title"]
    assert len(visualization.build_table(result)["rows"]) == 40


def test_sem_dado_nao_monta_grafico():
    assert visualization.build_visualizations(_result("hours_by_client", "client", [])) == []
    assert visualization.build_visualizations(_result("total_hours", None, [], total=None)) == []


def test_valor_nulo_fica_fora_do_grafico():
    result = _result("generation_time", "format", [{"label": "PDF", "value": None}, {"label": "XLSX", "value": 300.0}], unit="ms")
    assert visualization.build_visualizations(result)[0]["categories"] == ["XLSX"]


# --- SQL do reports_db (MySQL real) -----------------------------------------


def _payload(number: str, hours: float) -> dict:
    return {
        "packages": [{
            "header": {
                "project_code": number, "project_name": f"Projeto {number}",
                "location_date": "São Paulo, 01/01/2026", "month_label": "Setembro/2026",
                "signer1_name": "A", "signer1_company": "B", "signer2_name": "C", "signer2_company": "D",
            },
            "groups": [{"name": "ENG", "performance": 100.0, "activities": [{"description": "x", "hours": hours}]}],
        }],
        "formats": ["xlsx", "pdf"],
        "include_performance": False,
    }


@pytest.mark.reports_db
def test_repository_conta_relatorio_uma_vez_mesmo_com_dois_formatos(reports_db_engine, monkeypatch):
    monkeypatch.setattr(management, "MANAGEMENT_PANEL_LOGINS", {"gerente"})
    app.dependency_overrides[require_session] = lambda: {"name": "G", "login": "gerente", "email": "g@x"}
    try:
        with TestClient(app) as client:
            assert client.post("/generate", json=_payload("SE.CHAT.001", 8.0)).status_code == 200
            assert client.post("/generate", json=_payload("SE.CHAT.002", 3.0)).status_code == 200
    finally:
        app.dependency_overrides.pop(require_session, None)
    today = date.today()

    # 2 relatórios, 4 gerações (xlsx + pdf de cada) — conta relatório, não arquivo
    assert report_analytics_repository.count_generated_reports(today, today) == 2
    assert report_analytics_repository.generation_failures(today, today) == {"total": 4, "failed": 0}
    by_month = report_analytics_repository.generated_reports_by_month(today, today)
    assert by_month == [{"label": today.strftime("%Y-%m"), "value": 2}]
    hours = report_analytics_repository.report_hours_by_project(date(2026, 9, 1), date(2026, 9, 30), 10)
    assert hours == [{"label": "Projeto SE.CHAT.001", "value": 8.0}, {"label": "Projeto SE.CHAT.002", "value": 3.0}]


# --- exportação .xlsx ---------------------------------------------------------


def _export(client, **overrides):
    payload = {
        "title": "Horas por cliente — setembro/2026",
        "columns": ["Cliente", "Horas", "% do total"],
        "column_types": ["text", "hours", "percent"],
        "rows": [["ACME", 10.0, 76.9], ["=HYPERLINK(\"http://x\")", 3.0, 23.1]],
        "totals": ["Total", 13.0, 100.0],
        **overrides,
    }
    return client.post("/analytics/chat/export", json=payload)


def test_exporta_tabela_com_numeros_como_numeros_e_texto_sem_formula(monkeypatch):
    import io

    from openpyxl import load_workbook

    monkeypatch.setattr(management, "MANAGEMENT_PANEL_LOGINS", {"gerente"})
    app.dependency_overrides[require_session] = lambda: {"name": "G", "login": "gerente", "email": "g@x"}
    try:
        response = _export(TestClient(app))
    finally:
        app.dependency_overrides.pop(require_session, None)
    assert response.status_code == 200
    assert "horas-por-cliente-setembro-2026.xlsx" in response.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in ws[3]] == ["Cliente", "Horas", "% do total"]
    assert [c.value for c in ws[4]] == ["ACME", 10.0, 76.9]
    assert ws["B4"].number_format == '#,##0.00" h"'
    assert ws["A5"].data_type == "s"  # nome que parece fórmula continua texto
    assert [c.value for c in ws[6]] == ["Total", 13.0, 100.0]


def test_exportacao_recusa_numero_invalido_e_colaborador(monkeypatch):
    monkeypatch.setattr(management, "MANAGEMENT_PANEL_LOGINS", {"gerente"})
    app.dependency_overrides[require_session] = lambda: {"name": "G", "login": "gerente", "email": "g@x"}
    try:
        client = TestClient(app)
        assert client.post(
            "/analytics/chat/export",
            content='{"title": "t", "columns": ["a"], "rows": [[NaN]]}',
            headers={"content-type": "application/json"},
        ).status_code == 422
        app.dependency_overrides[require_session] = lambda: {"name": "C", "login": "colab", "email": "c@x"}
        assert _export(client).status_code == 403
    finally:
        app.dependency_overrides.pop(require_session, None)
