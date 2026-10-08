"""Horas externas: planilha própria de colaboradores que não apontam no Projectile
(`external_hours.py`, `POST /parse-external`, `GET /parse-external/template`).
Planilhas montadas em memória — nenhum teste depende de arquivo real."""

from __future__ import annotations

import inspect
import io
from datetime import date, datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from backend.app import external_hours
from backend.app.api.routers import external_hours as router_module
from backend.app.api.routers import parsing as parsing_module
from backend.app.core import authz
from backend.app.main import app, require_session
from backend.app.services import audit

HEADER = ["Data", "Colaborador", "Projeto", "Pacote de Trabalho", "Descrição", "Horas"]


def _xlsx(rows: list[list], header: list | None = HEADER, lead: int = 0) -> bytes:
    wb = Workbook()
    ws = wb.active
    for _ in range(lead):
        ws.append([])
    if header is not None:
        ws.append(header)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse(rows, **kw):
    return external_hours.parse_external_hours(io.BytesIO(_xlsx(rows, **kw)))


def test_le_linha_valida_com_data_texto_e_datetime():
    rows, issues = _parse(
        [
            [datetime(2026, 8, 12), "Fulano", "Projeto A 08.2026", "Pacote 1", "Revisão do desenho", 3.5],
            ["13/08/2026", "Beltrana", "Projeto A 08.2026", "Pacote 1", "Modelagem", "2,5"],
        ]
    )
    assert issues == []
    assert [r.date for r in rows] == [date(2026, 8, 12), date(2026, 8, 13)]
    assert [r.hours for r in rows] == [3.5, 2.5]
    assert rows[0].collaborator == "Fulano" and rows[0].package == "Pacote 1"
    assert rows[0].row == 2  # linha da planilha (cabeçalho é a 1)


def test_aceita_variacoes_de_cabecalho_e_cabecalho_fora_da_primeira_linha():
    header = ["dia", "PESSOA", "projeto", "Pacote", "Atividade", "Hs"]
    rows, issues = _parse([["12/08/2026", "Fulano", "P", "K", "Algo", 1]], header=header, lead=3)
    assert issues == [] and len(rows) == 1
    assert rows[0].row == 5  # 3 linhas em branco + cabeçalho


def test_coluna_faltando_diz_quais_faltaram():
    with pytest.raises(ValueError, match="faltam as colunas.*Horas"):
        _parse([["12/08/2026", "F", "P", "K", "D"]], header=["Data", "Colaborador", "Projeto", "Pacote de Trabalho", "Descrição"])


def test_planilha_sem_cabecalho_do_modelo_e_recusada():
    with pytest.raises(ValueError, match="cabeçalho do modelo"):
        _parse([["a", "b"]], header=["x", "y"])


def test_linha_em_branco_e_subtotal_sao_silenciosos():
    rows, issues = _parse(
        [
            ["12/08/2026", "Fulano", "P", "K", "Algo", 2],
            [None, None, None, None, None, None],
            [None, None, None, None, None, 2],  # subtotal solto
        ]
    )
    assert len(rows) == 1 and issues == []


@pytest.mark.parametrize(
    ("row", "reason"),
    [
        (["12/08/2026", "", "P", "K", "Algo", 1], "campo_vazio"),
        (["12/08/2026", "F", "P", "K", "", 1], "campo_vazio"),
        (["12-ago", "F", "P", "K", "Algo", 1], "data_invalida"),
        ([None, "F", "P", "K", "Algo", 1], "data_invalida"),
        (["12/08/2026", "F", "P", "K", "Algo", "abc"], "horas_invalidas"),
        (["12/08/2026", "F", "P", "K", "Algo", 0], "horas_invalidas"),
        (["12/08/2026", "F", "P", "K", "Algo", -2], "horas_invalidas"),
        (["12/08/2026", "F", "P", "K", "Algo", 25], "horas_invalidas"),
        (["12/08/2026", "F", "P", "K", "x" * 501, 1], "texto_longo"),
    ],
)
def test_linha_invalida_vira_aviso_e_fica_de_fora(row, reason):
    rows, issues = _parse([["12/08/2026", "Ok", "P", "K", "Boa", 1], row])
    assert [r.collaborator for r in rows] == ["Ok"]
    assert [i.reason for i in issues] == [reason]
    assert issues[0].row == 3 and "Linha 3" in issues[0].message


def test_numero_de_serie_do_excel_vira_data():
    rows, _ = _parse([[46000, "F", "P", "K", "Algo", 1]])  # 46000 = 2025-12-09 no calendário do Excel
    assert rows[0].date == date(2025, 12, 9)


def test_texto_e_limpo_de_controle_e_espaco_duplo():
    rows, _ = _parse([["12/08/2026", "  Fulano   de   Tal ", "P", "K", "Revisão\n do   desenho", 1]])
    assert rows[0].collaborator == "Fulano de Tal"
    assert rows[0].description == "Revisão do desenho"
    # o openpyxl nem grava caractere de controle, mas um arquivo vindo de fora pode trazê-lo
    assert external_hours._clean("a\x07b\x00 c") == "ab c"


def test_limite_de_linhas(monkeypatch):
    monkeypatch.setattr(external_hours, "MAX_ROWS", 3)
    ok = [["12/08/2026", "F", "P", "K", f"d{i}", 1] for i in range(3)]
    assert len(_parse(ok)[0]) == 3
    with pytest.raises(ValueError, match="mais de 3 linhas"):
        _parse(ok + [["12/08/2026", "F", "P", "K", "d3", 1]])


def test_modelo_gerado_e_lido_pelo_proprio_parser_e_nao_traz_exemplo_na_aba_de_lancamentos():
    content = external_hours.build_template()
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Horas externas", "Instruções"]
    rows, issues = external_hours.parse_external_hours(io.BytesIO(content))
    assert rows == [] and issues == []  # exemplo mora só na aba de instruções


# --- rotas ------------------------------------------------------------------


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr(authz, "MANAGEMENT_PANEL_LOGINS", {"gerente"})
    monkeypatch.setattr(authz, "COORDINATOR_LOGINS", {"coord"})
    events = []
    monkeypatch.setattr(audit, "record_event", lambda **kw: events.append(kw))
    state = {"user": {"name": "Gerente", "login": "gerente", "email": "g@x"}, "events": events}
    app.dependency_overrides[require_session] = lambda: state["user"]
    yield TestClient(app), state
    app.dependency_overrides.pop(require_session, None)


def _post(client, content: bytes):
    return client.post("/parse-external", files={"file": ("horas.xlsx", content, "application/octet-stream")})


def test_rota_devolve_linhas_e_avisos_e_audita_so_contagens(client):
    http, state = client
    body = _post(http, _xlsx([["12/08/2026", "Fulano", "P", "K", "Segredo do cliente", 2], ["xx", "F", "P", "K", "D", 1]]))
    assert body.status_code == 200
    data = body.json()
    assert data["rows"] == [
        {"row": 2, "date": "2026-08-12", "collaborator": "Fulano", "project": "P", "package": "K", "description": "Segredo do cliente", "hours": 2.0}
    ]
    assert [i["reason"] for i in data["issues"]] == ["data_invalida"]
    [event] = state["events"]
    assert event["action"] == "external_hours_parsed"
    assert event["metadata"] == {"rows": 1, "issues": 1, "hours": 2.0}
    assert "Segredo" not in str(event)


def test_coordenador_pode_e_colaborador_nao(client):
    http, state = client
    content = _xlsx([["12/08/2026", "F", "P", "K", "D", 1]])
    state["user"] = {"name": "Coord", "login": "coord", "email": "c@x"}
    assert _post(http, content).status_code == 200
    assert http.get("/parse-external/template").status_code == 200
    state["user"] = {"name": "Colab", "login": "colab", "email": "x@x"}
    assert _post(http, content).status_code == 403
    assert http.get("/parse-external/template").status_code == 403


def test_arquivo_invalido_e_planilha_fora_do_modelo_dao_400(client):
    http, _ = client
    assert _post(http, b"isto nao e um xlsx").status_code == 400
    bad = _post(http, _xlsx([["a", "b"]], header=["x", "y"]))
    assert bad.status_code == 400 and "cabeçalho do modelo" in bad.json()["detail"]


def test_upload_maior_que_o_limite_da_413_e_nao_deixa_temporario(client, monkeypatch):
    http, _ = client
    monkeypatch.setattr(parsing_module, "_MAX_UPLOAD_BYTES", 1024)
    assert _post(http, b"x" * 4096).status_code == 413


def test_rota_de_upload_leva_a_leitura_bloqueante_pro_threadpool():
    source = inspect.getsource(router_module.parse_external_endpoint)
    assert "run_in_threadpool(_parse_upload" in source
    assert not inspect.iscoroutinefunction(router_module.external_hours_template)


def test_modelo_baixa_como_xlsx(client):
    http, _ = client
    response = http.get("/parse-external/template")
    assert response.status_code == 200
    assert "modelo-horas-externas.xlsx" in response.headers["content-disposition"]
    assert load_workbook(io.BytesIO(response.content)).sheetnames[0] == "Horas externas"
